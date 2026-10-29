# This module will be responsible for parsing the data into the
# correct objects from the given data. This way, it can be reused with
# more data, rather then hard coding the values, allowing for scalability.

# Using load_workbook from openpyxl to parse data.
from openpyxl import load_workbook
from package import Package
from hash_table import HashTable
from datetime import datetime, time

# Loads the package data into a custom hash table, deals with special
# notes explicitly with a helper.
def load_packages(file_path):
    workbook = load_workbook(file_path)
    sheet = workbook["Sheet1"]
    rows = list(sheet.iter_rows(
    min_row=9,
    max_row=48,
    min_col=1,
    max_col=8,
    values_only=True
))
    packages = HashTable(len(rows))

    for row in rows:
        package_id, address, city, state, zip_code, deadline, weight, special_note = row
        address = _normalize_address(address)
        package = Package(
            package_id,
            address,
            city,
            state,
            str(zip_code),
            deadline,
            weight,
            special_note
    )
        
        _parse_special_notes(package)
        _parse_deadline(package)
        packages.insert(package.id, package)

    return packages
        

def _parse_special_notes(package):
    if not package.special_note:
        return
    # Can only be on X truck, detected by the string "Can only be on truck",
    # and must have the truck number as the last word in the string.
    elif "Can only be on truck" in package.special_note:
        parts = package.special_note.split()
        truck_number = int(parts[-1])
        package.required_truck = truck_number

    # Package delayed is detected by the string "Delayed" some where in the note
    # and it must end with "until hh:mm am/pm" to be valid.
    elif "Delayed" in package.special_note:
        time_string = package.special_note.split("until ")[1]
        package.available_at = _time_to_minutes(time_string)

    # Must be delivered with x,y,...,z detected by the string "Must be delivered with".
    # Extracts the package IDs listed after "with", separated by commas
    elif "Must be delivered with" in package.special_note:
        delivered_with = package.special_note.split("with ")[1]
        ids = delivered_with.split(",")
        grouped_ids = [int(package_id.strip()) for package_id in ids]
        package.grouped_with = grouped_ids

    elif "Wrong address" in package.special_note:
        package.address_is_valid = False


# Use this helper to convert any given time 12 hour in the format of HH:MM am/pm.
def _time_to_minutes(time_value):
    if isinstance(time_value, time):
        return time_value.hour * 60 + time_value.minute

    parsed_time = datetime.strptime(time_value.strip(), "%I:%M %p")
    return parsed_time.hour * 60 + parsed_time.minute

# This returns two objects, one containing the mapping of addresses to indices.
# The other is a distance matrix that takes two indices and gives you back the distance.
def load_distance_data(file_path):
    address_count = 0
    address_mapping = {}
    workbook = load_workbook(file_path)
    sheet = workbook["Sheet1"]
    rows = sheet.iter_rows(
        min_row=9,
        max_row=35,
        min_col=2,
        max_col=2,
        values_only=True
)
    for row in rows:
        address = row[0].split("\n")[0].strip()
        address = _normalize_address(address)
        address_mapping[address] = address_count
        address_count += 1

    distance_matrix = [[None] * address_count for _ in range(address_count)]
    rows2 = sheet.iter_rows(
        min_row=9,
        max_row=35,
        min_col=3,
        max_col=29,
        values_only=True
)
    for row_index, row in enumerate(rows2):
        for column_index, distance in enumerate(row):
            if distance is not None:
                distance = round(distance, 1)
                # Since A -> B is the same distance as B -> A, we
                # populate both directions each pass.
                distance_matrix[row_index][column_index] = distance
                distance_matrix[column_index][row_index] = distance

    return address_mapping, distance_matrix

# To get consistency with deadline types, we convert everything
# to minutes after midnight, including EOD.
def _parse_deadline(package):
    if package.deadline == "EOD":
        package.deadline = 1440
    else:
        package.deadline = _time_to_minutes(package.deadline)

# During main call, address on the packages and parsed data from the
# delivery locations where not equal strings. Parsed data used N/E/S/W while
# packages used North/East/South/West on addresses. This is used to convert.
def _normalize_address(address):
    return (
        address
        .replace(" South ", " S ")
        .replace(" North ", " N ")
        .replace(" East ", " E ")
        .replace(" West ", " W ")
        .strip()
    )